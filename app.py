# app.py
import os
import tempfile
import streamlit as st
from docx import Document as DocxDocument
from pypdf import PdfReader
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_community.vectorstores import FAISS
from langchain_groq import ChatGroq
from langchain_core.prompts import ChatPromptTemplate
from langchain_core.runnables import RunnablePassthrough
from langchain_core.output_parsers import StrOutputParser

st.set_page_config(
    page_title="Multi-Doc RAG Workspace",
    page_icon="🤖",
    layout="wide",
    initial_sidebar_state="expanded"
)

st.markdown("""
<style>
    /* Dark Theme & Glassmorphism Styling */
    .stApp {
        background: linear-gradient(135deg, #0f172a 0%, #1e293b 100%);
        color: #f8fafc;
    }
    [data-testid="stSidebar"] {
        background-color: rgba(30, 41, 59, 0.7) !important;
        backdrop-filter: blur(12px);
        border-right: 1px solid rgba(255, 255, 255, 0.1);
    }
    .stChatMessage {
        background-color: rgba(51, 65, 85, 0.5) !important;
        border-radius: 12px;
        border: 1px solid rgba(255, 255, 255, 0.05);
        margin-bottom: 12px;
    }
    .stButton>button {
        border-radius: 8px;
        background-color: #3b82f6;
        color: white;
        border: none;
        transition: all 0.2s ease-in-out;
    }
    .stButton>button:hover {
        background-color: #2563eb;
        box-shadow: 0 4px 12px rgba(59, 130, 246, 0.3);
    }
    .source-box {
        font-size: 0.85rem;
        padding: 8px 12px;
        background-color: rgba(15, 23, 42, 0.6);
        border-left: 3px solid #3b82f6;
        border-radius: 4px;
        margin-top: 6px;
    }
</style>
""", unsafe_allow_class_name=True)

# ------------------------------------------------------------------------------
# Document Extractors
# ------------------------------------------------------------------------------
def extract_text_from_file(uploaded_file) -> list[Document]:
    documents = []
    file_name = uploaded_file.name
    file_ext = os.path.splitext(file_name)[-1].lower()

    if file_ext == ".pdf":
        pdf_reader = PdfReader(uploaded_file)
        for i, page in enumerate(pdf_reader.pages):
            text = page.extract_text()
            if text and text.strip():
                documents.append(
                    Document(
                        page_content=text,
                        metadata={"source": file_name, "page": i + 1}
                    )
                )

    elif file_ext == ".docx":
        doc = DocxDocument(uploaded_file)
        full_text = []
        for p in doc.paragraphs:
            if p.text.strip():
                full_text.append(p.text)
        if full_text:
            documents.append(
                Document(
                    page_content="\n".join(full_text),
                    metadata={"source": file_name, "page": 1}
                )
            )

    elif file_ext == ".txt":
        content = uploaded_file.read().decode("utf-8")
        if content.strip():
            documents.append(
                Document(
                    page_content=content,
                    metadata={"source": file_name, "page": 1}
                )
            )

    return documents

# ------------------------------------------------------------------------------
# Caching Functions
# ------------------------------------------------------------------------------
@st.cache_resource(show_spinner=False)
def load_embeddings():
    return HuggingFaceEmbeddings(
        model_name="sentence-transformers/all-MiniLM-L6-v2",
        model_kwargs={'device': 'cpu'},
        encode_kwargs={'normalize_embeddings': True}
    )

def create_vector_store(documents):
    text_splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200)
    splits = text_splitter.split_documents(documents)
    embeddings = load_embeddings()
    vectorstore = FAISS.from_documents(splits, embeddings)
    return vectorstore

# ------------------------------------------------------------------------------
# Sidebar Initialization
# ------------------------------------------------------------------------------
with st.sidebar:
    st.title("⚙️ RAG Settings")
    
    api_key_input = st.text_input("Groq API Key", type="password", help="Enter your GROQ_API_KEY")
    groq_api_key = api_key_input or os.environ.get("GROQ_API_KEY", "")

    st.markdown("---")
    st.subheader("📁 Upload Documents")
    uploaded_files = st.file_uploader(
        "Upload PDF, TXT, or DOCX files",
        type=["pdf", "docx", "txt"],
        accept_multiple_files=True
    )

    if st.button("Process Documents", use_container_width=True):
        if not groq_api_key:
            st.error("Please provide a Groq API Key.")
        elif not uploaded_files:
            st.warning("Please upload at least one file.")
        else:
            with st.spinner("Processing files and building index..."):
                all_docs = []
                for uploaded_file in uploaded_files:
                    docs = extract_text_from_file(uploaded_file)
                    all_docs.extend(docs)

                if all_docs:
                    st.session_state.vectorstore = create_vector_store(all_docs)
                    st.success(f"Successfully processed {len(uploaded_files)} file(s)!")
                else:
                    st.error("Could not extract text from uploaded files.")

    st.markdown("---")
    if st.button("Clear Chat", use_container_width=True):
        st.session_state.messages = []
        st.rerun()

# ------------------------------------------------------------------------------
# Main Layout & Chat Interface
# ------------------------------------------------------------------------------
st.title("🤖 Multi-Document RAG Assistant")
st.caption("Powered by Groq `openai/gpt-oss-120b`, FAISS, and LangChain")

if "messages" not in st.session_state:
    st.session_state.messages = []

for message in st.session_state.messages:
    with st.chat_message(message["role"]):
        st.markdown(message["content"])
        if "sources" in message and message["sources"]:
            with st.expander("📚 View Retrieved Sources"):
                for src in message["sources"]:
                    st.markdown(
                        f"""<div class="source-box">
                        <b>File:</b> {src['source']} (Page {src['page']})<br>
                        <i>"{src['snippet']}..."</i>
                        </div>""",
                        unsafe_allow_html=True
                    )

user_query = st.chat_input("Ask a question about your uploaded documents...")

if user_query:
    st.session_state.messages.append({"role": "user", "content": user_query})
    with st.chat_message("user"):
        st.markdown(user_query)

    if "vectorstore" not in st.session_state or st.session_state.vectorstore is None:
        with st.chat_message("assistant"):
            st.warning("Please upload and process documents first before asking questions.")
    elif not groq_api_key:
        with st.chat_message("assistant"):
            st.error("Missing Groq API key. Provide it in the sidebar.")
    else:
        with st.chat_message("assistant"):
            with st.spinner("Searching documents & generating answer..."):
                retriever = st.session_state.vectorstore.as_retriever(search_kwargs={"k": 4})
                retrieved_docs = retriever.invoke(user_query)

                sources = [
                    {
                        "source": doc.metadata.get("source", "Unknown"),
                        "page": doc.metadata.get("page", 1),
                        "snippet": doc.page_content[:200].replace("\n", " ")
                    }
                    for doc in retrieved_docs
                ]

                context = "\n\n".join([doc.page_content for doc in retrieved_docs])
                prompt = ChatPromptTemplate.from_template(
                    """You are an expert AI assistant. Answer the user's question based ONLY on the context provided below.
If the information is not contained within the context, state clearly that you cannot find the answer in the provided documents.

Context:
{context}

Question:
{question}

Answer:"""
                )

                llm = ChatGroq(
                    groq_api_key=groq_api_key,
                    model_name="openai/gpt-oss-120b",
                    temperature=0.1
                )

                chain = (
                    {"context": lambda x: context, "question": RunnablePassthrough()}
                    | prompt
                    | llm
                    | StrOutputParser()
                )

                response_text = chain.invoke(user_query)
                st.markdown(response_text)

                if sources:
                    with st.expander("📚 View Retrieved Sources"):
                        for src in sources:
                            st.markdown(
                                f"""<div class="source-box">
                                <b>File:</b> {src['source']} (Page {src['page']})<br>
                                <i>"{src['snippet']}..."</i>
                                </div>""",
                                unsafe_allow_html=True
                            )

                st.session_state.messages.append({
                    "role": "assistant",
                    "content": response_text,
                    "sources": sources
                })
